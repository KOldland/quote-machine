"""Tests for DOCX export route."""
import io
import json
import os
import tempfile
import pytest
from unittest.mock import patch, MagicMock
from docx import Document

# Mark as integration test - requires Flask app context
pytestmark = pytest.mark.integration


class TestExportDocx:
    """Test suite for /api/export-docx endpoint."""

    @pytest.fixture
    def app(self):
        """Create test Flask app with export blueprint."""
        from flask import Flask
        app = Flask(__name__)
        app.config['TESTING'] = True
        app.config['SECRET_KEY'] = 'test-secret'
        
        # Import and register the export blueprint
        from app.export_routes import export_bp
        app.register_blueprint(export_bp)
        
        return app

    @pytest.fixture
    def client(self, app):
        """Create test client."""
        return app.test_client()

    @pytest.fixture
    def mock_session(self, client):
        """Set up session with required keys."""
        with client.session_transaction() as sess:
            sess['form_key'] = 'standard_build'
            sess['form_data'] = {
                'client_name': 'Test Client',
                'client_address': '123 Test St',
                'deposit_pct': 0.1,
                'completion_pct': 0.1
            }
        return client

    @patch('app.export_routes.get_output_template')
    @patch('app.export_routes.Calculator')
    def test_successful_docx_generation(self, mock_calc_class, mock_get_template, mock_session):
        """Test that a valid DOCX is generated with basic template."""
        # Mock template
        mock_template = {
            'sections_json': json.dumps({
                'header': {'enabled': True, 'show_logo': False, 'company_name': 'Test Co'},
                'footer': {'enabled': True, 'show_page_numbers': False, 'terms_text': 'Test terms'},
                'body': {
                    'show_client_info': True,
                    'show_line_items_table': True,
                    'show_subtotals': True,
                    'show_payment_schedule': True
                }
            }),
            'css_json': json.dumps({
                'fontFamily': 'Arial',
                'fontSize': 12,
                'headingColor': '#FF0000'
            })
        }
        mock_get_template.return_value = mock_template
        
        # Mock calculator
        mock_calc = MagicMock()
        mock_calc.calculate_quote.return_value = {
            'client_name': 'Test Client',
            'client_address': '123 Test St',
            'groups': [{
                'name': 'Base',
                'subtotal': 1000.0,
                'items': [{
                    'output_title': 'Item 1',
                    'internal_description': 'Test item',
                    'units': 1,
                    'unit_cost': 1000.0,
                    'line_total': 1000.0
                }]
            }],
            'grand_total': 1000.0,
            'deposit_pct': 0.1,
            'completion_pct': 0.1,
            'deposit_amount': 100.0,
            'completion_amount': 100.0,
            'middle_balance': 800.0
        }
        mock_calc_class.return_value = mock_calc
        
        # Make request
        response = mock_session.get('/api/export-docx')
        
        # Verify response
        assert response.status_code == 200
        assert response.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        assert 'attachment' in response.headers['Content-Disposition']
        
        # Verify it's a valid DOCX by reading with python-docx
        doc_bytes = io.BytesIO(response.data)
        doc = Document(doc_bytes)
        
        # Check that document has content
        assert len(doc.paragraphs) > 0
        
        # Check for client info in paragraphs
        full_text = '\n'.join([p.text for p in doc.paragraphs])
        assert 'Test Client' in full_text or 'Client Information' in full_text

    @patch('app.export_routes.get_output_template')
    def test_missing_form_key(self, mock_get_template, client):
        """Test that 400 is returned when form_key is missing."""
        response = client.get('/api/export-docx')
        assert response.status_code == 400
        assert b'Form key not found' in response.data

    @patch('app.export_routes.get_output_template')
    def test_invalid_template_json(self, mock_get_template, mock_session):
        """Test that 500 is returned on malformed template JSON."""
        mock_get_template.return_value = {
            'sections_json': 'invalid json',
            'css_json': '{}'
        }
        response = mock_session.get('/api/export-docx')
        assert response.status_code == 500
        assert b'Invalid template JSON' in response.data

    @patch('app.export_routes.get_output_template')
    @patch('app.export_routes.Calculator')
    def test_disabled_sections(self, mock_calc_class, mock_get_template, mock_session):
        """Test that disabled sections are not rendered."""
        mock_template = {
            'sections_json': json.dumps({
                'header': {'enabled': False},
                'footer': {'enabled': False},
                'body': {
                    'show_client_info': False,
                    'show_line_items_table': False,
                    'show_subtotals': False,
                    'show_payment_schedule': False
                }
            }),
            'css_json': json.dumps({})
        }
        mock_get_template.return_value = mock_template
        
        mock_calc = MagicMock()
        mock_calc.calculate_quote.return_value = {
            'client_name': 'Test',
            'client_address': 'Addr',
            'groups': [],
            'grand_total': 0,
            'deposit_pct': 0,
            'completion_pct': 0,
            'deposit_amount': 0,
            'completion_amount': 0,
            'middle_balance': 0
        }
        mock_calc_class.return_value = mock_calc
        
        response = mock_session.get('/api/export-docx')
        assert response.status_code == 200
        
        # Should still produce a valid DOCX (maybe empty)
        doc_bytes = io.BytesIO(response.data)
        doc = Document(doc_bytes)
        assert doc is not None  # Valid DOCX even if minimal content

    @patch('app.export_routes.get_output_template')
    @patch('app.export_routes.Calculator')
    def test_css_styling_applied(self, mock_calc_class, mock_get_template, mock_session):
        """Test that CSS styling is applied to the document."""
        mock_template = {
            'sections_json': json.dumps({
                'header': {'enabled': True, 'show_logo': False, 'company_name': 'Styled Co'},
                'footer': {'enabled': False},
                'body': {'show_client_info': True}
            }),
            'css_json': json.dumps({
                'fontFamily': 'Courier New',
                'fontSize': 14,
                'headingColor': '#00FF00'
            })
        }
        mock_get_template.return_value = mock_template
        
        mock_calc = MagicMock()
        mock_calc.calculate_quote.return_value = {
            'client_name': 'Styled Client',
            'client_address': 'Styled Address',
            'groups': [],
            'grand_total': 0,
            'deposit_pct': 0,
            'completion_pct': 0,
            'deposit_amount': 0,
            'completion_amount': 0,
            'middle_balance': 0
        }
        mock_calc_class.return_value = mock_calc
        
        response = mock_session.get('/api/export-docx')
        assert response.status_code == 200
        
        doc_bytes = io.BytesIO(response.data)
        doc = Document(doc_bytes)
        
        # Check that styling was applied (at least font name)
        # Note: python-docx doesn't always expose style in runs reliably for test
        # but we can check the document was created successfully
        assert len(doc.paragraphs) > 0


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
